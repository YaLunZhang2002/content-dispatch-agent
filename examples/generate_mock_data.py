# -*- coding: utf-8 -*-
"""
Generate synthetic (desensitized) Excel dataset for testing and open-source demonstration.
No real customer confidential text or personal info is included.
"""

import os
import openpyxl


def generate_synthetic_workbook(output_path: str = "examples/sample_pool.xlsx"):
    wb = openpyxl.Workbook()

    # Sheet 1: 今日头条 (带标题、正文)
    ws_tt = wb.active
    ws_tt.title = "今日头条"
    ws_tt.append(["标题", "正文", "话题标签"])
    for i in range(1, 41):
        ws_tt.append([
            f"大屏科技新标杆：探索第{i}代高刷纯平画质之美",
            f"居家客厅换新指南第{i}期！在这个时代，好的电视不仅是一块屏幕，更是家庭艺术的一部分。\n"
            f"纯平贴墙设计让电线隐形，色彩调校细腻温和，告别眩光与反光，深夜追剧护眼更沉浸。",
            "#科技好物 #家电测评 #客厅改造"
        ])

    # Sheet 2: 小红书 (短篇、重生活感、丰富 Emoji)
    ws_xhs = wb.create_sheet("小红书")
    ws_xhs.append(["标题", "正文", "话题标签"])
    for i in range(1, 41):
        ws_xhs.append([
            f"✨租房党梦中情屏！第{i}套极简客厅软装灵感",
            f"拆箱惊艳到全家！真的薄得像一幅画🖼️，直接上墙毫无压迫感～\n"
            f"和闺蜜窝在沙发上刷剧超有氛围感，开灯不反光，色彩还原太真实了，大爱！",
            "#家居美学 #小户型装修 #壁纸电视 #生活碎片"
        ])

    # Sheet 3: 抖音短视频脚本 (无显式标题，重口语化)
    ws_dy = wb.create_sheet("抖音")
    ws_dy.append(["标题", "正文", "话题标签"])
    for i in range(1, 31):
        ws_dy.append([
            "",
            f"很多人买电视都看错参数了！听劝，先看这三点！第{i}个细节就是屏幕本身的物理控色能力。\n"
            f"不仅白天不泛白，动态拖影也能直接拉满。带你实测这台最新旗舰大屏！",
            "#电视选购 #家电避坑 #硬核数码"
        ])

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)
    wb.close()
    return output_path


if __name__ == "__main__":
    generate_synthetic_workbook()
